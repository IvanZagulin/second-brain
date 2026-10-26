"""Bank statement adapters. Each returns a StatementData with normalised rows; the matching layer
does dedupe and merging. Supported: T-Bank CSV/XLSX, generic CSV/XLSX with a recognisable
header, 1C ClientBankExchange (business accounts), Sber PDF statements (best effort)."""

from __future__ import annotations

import csv
import hashlib
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, time
from pathlib import Path

from app.core.money import to_minor
from app.core.timeutil import tz


@dataclass
class StatementRow:
    occurred_at: datetime
    has_time: bool
    amount_minor: int
    direction: str
    description: str = ""
    merchant: str | None = None
    bank_category: str | None = None
    mcc: str | None = None
    card_last4: str | None = None
    bank_op_id: str | None = None
    balance_after_minor: int | None = None
    currency: str = "RUB"
    counterparty_account: str | None = None
    flow_hint: str | None = None
    scope_hint: str | None = None
    category_hint: str | None = None
    fingerprint: str = ""


@dataclass
class StatementData:
    bank: str
    format: str
    account_hint: str | None = None  # last digits of the account / card
    rows: list[StatementRow] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    closing_balance_minor: int | None = None  # balance confirmed by the bank at closing_at
    closing_at: datetime | None = None


class UnsupportedStatement(ValueError):
    pass


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


_DATE_FORMATS = ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S",
                 "%Y-%m-%d", "%d.%m.%y %H:%M", "%d.%m.%y", "%d/%m/%Y")


def _parse_dt(value, tz_name: str) -> tuple[datetime, bool] | None:
    if isinstance(value, datetime):
        dt = value
        has_time = dt.time() != time(0, 0)
    else:
        s = str(value or "").strip()
        if not s:
            return None
        dt = None
        for fmt in _DATE_FORMATS:
            try:
                dt = datetime.strptime(s, fmt)
                break
            except ValueError:
                continue
        if dt is None:
            return None
        has_time = "%H" in fmt
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=tz(tz_name))
    return dt, has_time


def _fingerprints(bank: str, account: str | None, rows: list[StatementRow]) -> None:
    """Stable per-row keys. Identical rows inside one file get an occurrence counter, so re-importing
    the same statement yields the same keys while two genuine identical purchases stay distinct."""
    seen: Counter = Counter()
    for r in rows:
        base = "|".join([
            bank, account or "", r.occurred_at.strftime("%Y-%m-%dT%H:%M" if r.has_time else "%Y-%m-%d"),
            str(r.amount_minor), r.direction, re.sub(r"\s+", " ", r.description.lower())[:120], r.bank_op_id or "",
        ])
        seen[base] += 1
        r.fingerprint = hashlib.sha256(f"{base}|{seen[base]}".encode()).hexdigest()


# ---------------------------------------------------------------- tabular (CSV / XLSX)

_HEADER_ALIASES = {
    "date": ("дата операции", "дата и время операции", "дата", "date", "дата проводки", "дата транзакции"),
    "amount": ("сумма операции", "сумма", "amount", "сумма в валюте счета", "сумма платежа", "сумма в рублях"),
    "income": ("приход", "поступление", "зачисление", "кредит", "credit"),
    "expense": ("расход", "списание", "дебет", "debit"),
    "description": ("описание", "назначение платежа", "назначение", "description", "детали операции",
                    "комментарий", "наименование операции"),
    "merchant": ("контрагент", "получатель", "магазин", "merchant", "плательщик / получатель"),
    "category": ("категория", "category"),
    "mcc": ("mcc",),
    "card": ("номер карты", "карта", "card"),
    "status": ("статус", "status"),
    "currency": ("валюта операции", "валюта", "currency"),
    "balance": ("остаток", "баланс", "balance"),
    "op_id": ("номер операции", "id операции", "идентификатор", "номер документа"),
}


def _map_header(header: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    norm = [str(h or "").strip().lower() for h in header]
    for key, aliases in _HEADER_ALIASES.items():
        for alias in aliases:
            for i, h in enumerate(norm):
                if i in mapping.values():
                    continue
                if h == alias or (len(alias) > 4 and h.startswith(alias)):
                    mapping.setdefault(key, i)
                    break
            if key in mapping:
                break
    return mapping


def _rows_from_table(table: list[list], tz_name: str, bank_guess: str) -> StatementData:
    header_idx, mapping = None, {}
    for i, row in enumerate(table[:30]):
        m = _map_header([str(c) if c is not None else "" for c in row])
        if "date" in m and ("amount" in m or ("income" in m and "expense" in m)):
            header_idx, mapping = i, m
            break
    if header_idx is None:
        raise UnsupportedStatement("Не найдена строка заголовков с датой и суммой")
    header_text = " ".join(str(c or "").lower() for c in table[header_idx])
    bank = "tbank" if ("кэшбэк" in header_text and "mcc" in header_text) else bank_guess
    data = StatementData(bank=bank, format="table")
    cards: Counter = Counter()

    def cell(row, key):
        idx = mapping.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    for row in table[header_idx + 1:]:
        if not row or all(c in (None, "") for c in row):
            continue
        status = str(cell(row, "status") or "").strip().upper()
        if status in ("FAILED", "ОТКЛОНЕНА", "ОТМЕНЕНА"):
            continue
        parsed_dt = _parse_dt(cell(row, "date"), tz_name)
        if not parsed_dt:
            continue
        try:
            if "amount" in mapping:
                raw_amount = cell(row, "amount")
                if raw_amount in (None, ""):
                    continue
                minor = to_minor(raw_amount if not isinstance(raw_amount, float) else str(raw_amount))
                direction = "out" if minor < 0 else "in"
                minor = abs(minor)
            else:
                inc, exp = cell(row, "income"), cell(row, "expense")
                if inc not in (None, "", 0, "0"):
                    minor, direction = abs(to_minor(str(inc))), "in"
                elif exp not in (None, "", 0, "0"):
                    minor, direction = abs(to_minor(str(exp))), "out"
                else:
                    continue
        except ValueError:
            continue
        if minor == 0:
            continue
        card = str(cell(row, "card") or "")
        card4 = re.sub(r"\D", "", card)[-4:] or None
        if card4:
            cards[card4] += 1
        balance = cell(row, "balance")
        try:
            balance_minor = to_minor(str(balance)) if balance not in (None, "") else None
        except ValueError:
            balance_minor = None
        desc = str(cell(row, "description") or "").strip()
        merchant = str(cell(row, "merchant") or "").strip() or (desc if bank == "tbank" else None)
        data.rows.append(StatementRow(
            occurred_at=parsed_dt[0], has_time=parsed_dt[1], amount_minor=minor, direction=direction,
            description=desc, merchant=merchant or None, bank_category=str(cell(row, "category") or "") or None,
            mcc=str(cell(row, "mcc") or "").strip() or None, card_last4=card4,
            bank_op_id=str(cell(row, "op_id") or "").strip() or None, balance_after_minor=balance_minor,
            currency=(str(cell(row, "currency") or "RUB").strip().upper() or "RUB")[:3],
        ))
    data.account_hint = cards.most_common(1)[0][0] if cards else None
    return data


def parse_csv(raw: bytes, tz_name: str, bank_guess: str = "other") -> StatementData:
    text = _decode(raw)
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    table = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    data = _rows_from_table(table, tz_name, bank_guess)
    data.format = "csv"
    return data


def parse_xlsx(raw: bytes, tz_name: str, bank_guess: str = "other") -> StatementData:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    last_error: Exception | None = None
    for ws in wb.worksheets:
        table = [list(r) for r in ws.iter_rows(values_only=True)]
        try:
            data = _rows_from_table(table, tz_name, bank_guess)
            data.format = "xlsx"
            return data
        except UnsupportedStatement as exc:
            last_error = exc
    raise UnsupportedStatement(str(last_error or "Пустая таблица"))


# ---------------------------------------------------------------- 1C ClientBankExchange

def parse_1c(raw: bytes, tz_name: str) -> StatementData:
    text = _decode(raw)
    if "1CClientBankExchange" not in text[:200]:
        raise UnsupportedStatement("Не формат 1C")
    our_account = None
    m = re.search(r"^РасчСчет=(\d+)", text, re.MULTILINE)
    if m:
        our_account = m.group(1)
    data = StatementData(bank="tbank_business", format="1c", account_hint=our_account[-4:] if our_account else None)
    for block in re.findall(r"СекцияДокумент=.*?КонецДокумента", text, re.DOTALL):
        fields = dict(re.findall(r"^([^=\r\n]+)=(.*)$", block, re.MULTILINE))
        dt = _parse_dt(fields.get("Дата") or fields.get("ДатаСписано") or fields.get("ДатаПоступило"), tz_name)
        if not dt or not fields.get("Сумма"):
            continue
        try:
            minor = abs(to_minor(fields["Сумма"]))
        except ValueError:
            continue
        payer = fields.get("ПлательщикСчет") or fields.get("ПлательщикРасчСчет")
        outgoing = our_account is not None and payer == our_account
        counterparty = fields.get("Получатель1") or fields.get("Получатель") if outgoing else \
            fields.get("Плательщик1") or fields.get("Плательщик")
        data.rows.append(StatementRow(
            occurred_at=dt[0], has_time=False, amount_minor=minor, direction="out" if outgoing else "in",
            description=fields.get("НазначениеПлатежа", "").strip(), merchant=(counterparty or "").strip() or None,
            bank_op_id=fields.get("Номер"),
            counterparty_account=fields.get("ПолучательСчет") if outgoing else fields.get("ПлательщикСчет"),
        ))
    return data


# ---------------------------------------------------------------- Sber PDF

_SBER_LINE = re.compile(
    r"^(?P<date>\d{2}\.\d{2}\.\d{4})\s+(?P<time>\d{2}:\d{2})\s+(?:(?P<code>\d{4,8})\s+)?(?P<cat>.+?)\s+"
    r"(?P<amount>[+]?\d{1,3}(?:[  ]\d{3})*,\d{2})\s+(?P<balance>\d{1,3}(?:[  ]\d{3})*,\d{2})$"
)
_SBER_DESC = re.compile(r"^(?P<date>\d{2}\.\d{2}\.\d{4})\s+(?P<desc>.+)$")


_RU_MONTHS = {m: i + 1 for i, m in enumerate(
    ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября", "ноября", "декабря"])}
_NUM_RU = r"\d{1,3}(?:[  ]\d{3})*,\d{2}"
_NUM_DOT = r"\d{1,3}(?:[  ]\d{3})*\.\d{2}"


def parse_sber_certificate(text: str, tz_name: str) -> StatementData:
    """«Справка по операции» from Sber Online: exactly one operation."""
    data = StatementData(bank="sber", format="sber_certificate")
    acc = re.search(r"сч[её]ту\s*[•*]{1,4}\s?(\d{4})", text)
    data.account_hint = acc.group(1) if acc else None
    when = re.search(r"(\d{1,2})\s+([а-я]+)\s+(\d{4})\s+в\s+(\d{1,2}):(\d{2})", text)
    amount = re.search(rf"Сумма в валюте сч[её]та.*?\n\s*({_NUM_RU})", text, re.S)
    if not when or not amount or when.group(2) not in _RU_MONTHS:
        raise UnsupportedStatement("Не удалось прочитать справку по операции")
    dt = datetime(int(when.group(3)), _RU_MONTHS[when.group(2)], int(when.group(1)), int(when.group(4)),
                  int(when.group(5)), tzinfo=tz(tz_name))
    desc = re.search(r"Описание\s+Статус операции\s*\n(.+?)\n(?:\((.+?)\)\s*\n)?", text)
    merchant = desc.group(1).strip() if desc else None
    if merchant:
        merchant = re.sub(r"\s+(Платёж|Платеж)\s+(выполнен|отклон[её]н).*$", "", merchant).strip()
    raw_merchant = desc.group(2).strip() if desc and desc.group(2) else None
    op_type = re.search(r"Тип операции\s*\n.*?\d{1,2}:\d{2}\s+(.+)$", text, re.M)
    auth = re.search(r"Код авторизации\s*\n\s*(\d+)", text)
    direction = "in" if "операция зачисления" in text.lower() else "out"
    data.rows.append(StatementRow(
        occurred_at=dt, has_time=True, amount_minor=to_minor(amount.group(1)), direction=direction,
        description=" ".join(x for x in (merchant, f"({raw_merchant})" if raw_merchant else None,
                                         op_type.group(1).strip() if op_type else None) if x),
        merchant=raw_merchant or merchant, bank_op_id=auth.group(1) if auth else None))
    return data


def parse_payment_receipt(text: str, tz_name: str) -> StatementData:
    """A paid-payment confirmation (ЮMoney and similar): date, amount, purpose, recipient."""
    date_m = re.search(r"Дата:\s*(\d{2}\.\d{2}\.\d{2,4})", text)
    sum_m = re.search(r"СУММА\s+(\d[\d  ]*(?:[.,]\d{1,2})?)\s*[PР₽]", text)
    if not date_m or not sum_m:
        raise UnsupportedStatement("Не удалось прочитать квитанцию")
    parsed_dt = _parse_dt(date_m.group(1), tz_name)
    purpose = re.search(r"НАЗНАЧЕНИЕ ПЛАТЕЖА\s+(.+)", text)
    recipient = re.search(r"ПОЛУЧАТЕЛЬ\s+(.+)", text)
    bank = re.search(r"Банк:\s*(.+)", text)
    ext = re.search(r"ВНЕШНИЙ НОМЕР ТРАНЗАКЦИИ\s+(\S+)", text)
    purpose_s = purpose.group(1).strip() if purpose else ""
    recipient_s = recipient.group(1).strip() if recipient else ""
    is_tax = bool(re.search(r"налог|фнс|казначейств", f"{purpose_s} {recipient_s}", re.I))
    data = StatementData(bank="yoomoney" if bank and "мани" in bank.group(1).lower() else "other", format="receipt")
    data.rows.append(StatementRow(
        occurred_at=parsed_dt[0], has_time=False, amount_minor=to_minor(sum_m.group(1)), direction="out",
        description=purpose_s, merchant=recipient_s or None, bank_op_id=ext.group(1) if ext else None,
        flow_hint="tax" if is_tax else None, scope_hint="business" if is_tax else None,
        category_hint="taxes" if is_tax else None))
    return data


def parse_tbank_business_pdf(pdf, text: str, tz_name: str) -> StatementData:
    """T-Bank current-account statement (ИП). Columns are located by header word positions; an
    operation is the block of lines from a dated line to the next one."""
    acc = re.search(r"Выписка по счету №\s*(\d{20})", text)
    data = StatementData(bank="tbank_business", format="tbank_pdf", account_hint=acc.group(1)[-4:] if acc else None)
    closing = re.search(r"Исходящий остаток:\s*(\d[\d  ]*\.\d{2})", text)
    formed = re.search(r"Дата формирования:\s*(\d{2}\.\d{2}\.\d{4})\s*/\s*(\d{2}:\d{2}:\d{2})", text)
    if closing and formed:
        data.closing_balance_minor = to_minor(closing.group(1))
        parsed_formed = _parse_dt(f"{formed.group(1)} {formed.group(2)}", tz_name)
        data.closing_at = parsed_formed[0] if parsed_formed else None
    cols: dict[str, float] = {}
    blocks: list[list[list[dict]]] = []
    current: list[list[dict]] | None = None
    for page in pdf.pages:
        words = page.extract_words()
        for w in words:
            if w["text"] in ("Дебет", "Кредит", "Корреспондент", "Содержание") and w["text"] not in cols:
                cols[w["text"]] = w["x0"]
        lines: dict[int, list[dict]] = {}
        for w in words:
            lines.setdefault(round(w["top"]), []).append(w)
        for top in sorted(lines):
            line = sorted(lines[top], key=lambda w: w["x0"])
            first = line[0]
            if first["x0"] < 100 and re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", first["text"]) and \
                    any(w["text"] == "Вид:" for w in line):
                current = [line]
                blocks.append(current)
            elif first["x0"] < 100 and first["text"] in ("Обороты", "Остаток", "Итого"):
                current = None
            elif current is not None and first["x0"] >= 130:
                current.append(line)
    if not {"Дебет", "Кредит", "Содержание", "Корреспондент"} <= cols.keys():
        raise UnsupportedStatement("Не найдены колонки выписки Т-Банка")
    debit_x, credit_x = cols["Дебет"], cols["Кредит"]
    content_x, corr_x = cols["Содержание"], cols["Корреспондент"]
    for block in blocks:
        head = block[0]
        dt = _parse_dt(head[0]["text"], tz_name)
        amount_words = [w for w in head if w["x0"] >= debit_x - 8]
        if not dt or not amount_words:
            continue
        try:
            minor = to_minor("".join(w["text"] for w in amount_words))
        except ValueError:
            continue
        direction = "in" if amount_words[0]["x0"] >= credit_x - 8 else "out"
        corr = " ".join(w["text"] for w in head if corr_x - 2 <= w["x0"] < content_x - 2)
        content = " ".join(w["text"] for line in block for w in line if content_x - 2 <= w["x0"] < debit_x - 8)
        doc_no = None
        for line in block:
            for i, w in enumerate(line):
                if w["text"] == "Номер:" and i + 1 < len(line):
                    doc_no = line[i + 1]["text"]
        op_date = re.search(r"Дата операции:\s*(\d{2}\.\d{2}\.\d{4})", content)
        occurred = _parse_dt(op_date.group(1), tz_name) if op_date else dt
        content_clean = re.sub(r"\s*Дата операции:\s*\d{2}\.\d{2}\.\d{4}", "", content).strip()
        card = re.search(r"\d{4}\.{3}(\d{4})\s+(.+?)(?:\.\s*Договор|\s+Договор|$)", content_clean)
        row = StatementRow(occurred_at=(occurred or dt)[0], has_time=False, amount_minor=minor, direction=direction,
                           description=content_clean, merchant=corr or None, bank_op_id=doc_no)
        low = content_clean.lower()
        if card:
            row.card_last4 = card.group(1)
            row.merchant = re.sub(r"\s+[A-Za-z]+\s+RUS?\.?$", "", card.group(2)).strip(" .") or card.group(2)
            row.flow_hint, row.scope_hint = "expense", "personal"
        elif "перевод собственных средств" in low or "между своими" in low:
            row.flow_hint, row.scope_hint, row.category_hint = "transfer", "business", "transfers"
        elif low.startswith("плата за") or "комисси" in low or ("кэп" in low and direction == "out"):
            row.flow_hint, row.scope_hint, row.category_hint = "expense", "business", "business"
        elif "налог" in low or "уфк" in low or "фнс" in low:
            row.flow_hint, row.scope_hint, row.category_hint = "tax", "business", "taxes"
        elif direction == "in" and re.search(r"оплата по сч|платеж по сч|платёж по сч|по договору|за услуги", low):
            row.flow_hint, row.scope_hint, row.category_hint = "income", "business", "business_income"
        data.rows.append(row)
    if not data.rows:
        raise UnsupportedStatement("В выписке Т-Банка не найдено операций")
    return data


def parse_pdf(raw: bytes, tz_name: str) -> StatementData:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        text = "\n".join((p.extract_text() or "") for p in pdf.pages)
        if "Справка по операции" in text and "Сбербанк" in text:
            return parse_sber_certificate(text, tz_name)
        if "ТБанк" in text and "Выписка по счету" in text and "Кредит" in text:
            return parse_tbank_business_pdf(pdf, text, tz_name)
        if re.search(r"^\s*ОПЛАЧЕНО", text, re.M) and "СУММА" in text:
            return parse_payment_receipt(text, tz_name)
    return parse_sber_pdf(raw, tz_name)


def parse_sber_pdf(raw: bytes, tz_name: str) -> StatementData:
    import pdfplumber

    lines: list[str] = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for page in pdf.pages:
            lines.extend((page.extract_text() or "").splitlines())
    data = StatementData(bank="sber", format="pdf")
    m = re.search(r"[*•]{2,4}\s?(\d{4})", " ".join(lines[:40]))
    if m:
        data.account_hint = m.group(1)
    pending: StatementRow | None = None
    for line in lines:
        line = line.strip()
        hit = _SBER_LINE.match(line)
        if hit:
            dt = _parse_dt(f"{hit['date']} {hit['time']}", tz_name)
            if not dt:
                continue
            amount_raw = hit["amount"]
            pending = StatementRow(
                occurred_at=dt[0], has_time=True, amount_minor=abs(to_minor(amount_raw)),
                direction="in" if amount_raw.startswith("+") else "out", description="",
                bank_category=hit["cat"].strip(), bank_op_id=hit["code"],
                balance_after_minor=to_minor(hit["balance"]),
            )
            data.rows.append(pending)
            continue
        if pending is not None and not pending.description:
            d = _SBER_DESC.match(line)
            if d:
                desc = d["desc"].strip()
                pending.description = desc
                pending.merchant = re.split(r"\.\s*Операция по", desc)[0].strip() or None
                card = re.search(r"\*{2,4}(\d{4})", desc)
                if card:
                    pending.card_last4 = card.group(1)
    if not data.rows:
        raise UnsupportedStatement("Не удалось распознать операции в PDF. Пришлите CSV/XLSX или образец выписки.")
    return data


def parse_statement(file_name: str, raw: bytes, tz_name: str, bank_hint: str | None = None) -> StatementData:
    ext = Path(file_name).suffix.lower()
    head = raw[:200]
    if b"1CClientBankExchange" in head or "1CClientBankExchange".encode("cp1251") in head:
        data = parse_1c(raw, tz_name)
    elif ext in (".xlsx", ".xlsm"):
        data = parse_xlsx(raw, tz_name, bank_hint or "other")
    elif ext in (".csv", ".txt", ".tsv"):
        data = parse_csv(raw, tz_name, bank_hint or "other")
    elif ext == ".pdf":
        data = parse_pdf(raw, tz_name)
    else:
        raise UnsupportedStatement(f"Формат {ext or 'без расширения'} пока не поддерживается")
    if bank_hint and data.bank == "other":
        data.bank = bank_hint
    _fingerprints(data.bank, data.account_hint, data.rows)
    return data
